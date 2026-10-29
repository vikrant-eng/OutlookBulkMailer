"""
Export module: write campaign history to CSV or Excel.
"""

import csv
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from modules.database import get_campaign, get_recipients


def export_csv(campaign_id: int, filepath: str) -> int:
    """Export campaign recipients to CSV. Returns row count."""
    campaign = get_campaign(campaign_id)
    recipients = get_recipients(campaign_id)
    if not campaign or not recipients:
        return 0

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["name", "email", "status", "error_msg", "sent_at"])
        writer.writeheader()
        writer.writerows(recipients)
    return len(recipients)


def export_excel(campaign_id: int, filepath: str) -> int:
    """Export campaign recipients to a styled Excel file. Returns row count."""
    campaign = get_campaign(campaign_id)
    recipients = get_recipients(campaign_id)
    if not campaign or not recipients:
        return 0

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    # ── Campaign summary ─────────────────────────────────
    ws["A1"] = "Campaign:"
    ws["B1"] = campaign["name"]
    ws["A2"] = "Subject:"
    ws["B2"] = campaign["subject"]
    ws["A3"] = "Exported:"
    ws["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws.append([])

    # ── Header row ───────────────────────────────────────
    header = ["Name", "Email", "Status", "Error", "Sent At"]
    ws.append(header)
    header_row = ws.max_row
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    for col, cell in enumerate(ws[header_row], start=1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # ── Data rows ────────────────────────────────────────
    status_colors = {"sent": "D4EDDA", "failed": "F8D7DA", "pending": "FFF3CD"}
    for r in recipients:
        ws.append([r["name"], r["email"], r["status"], r.get("error_msg", ""), r.get("sent_at", "")])
        row_cells = ws[ws.max_row]
        color = status_colors.get(r["status"], "FFFFFF")
        fill = PatternFill("solid", fgColor=color)
        for cell in row_cells:
            cell.fill = fill

    # ── Auto-fit columns ─────────────────────────────────
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    wb.save(filepath)
    return len(recipients)
